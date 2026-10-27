"""Turn the plate images and the building-block register into a print page and a Word file.

The print page (HTML, A4 landscape) is rendered to PDF by a headless browser; the Word file is
written with python-docx, which is a documentation tool only (``make diagrams`` adds it with
``uv run --with``; it is not a project dependency and never ships).
"""

import html
import re
import struct
from pathlib import Path
from typing import Any

from scripts.diagrams.content import (
    BLOCKS,
    INTRO,
    LEGEND_TEXT,
    PLATES,
    SUBTITLE,
    TITLE,
    Family,
)

FAMILY_HEX = {
    Family.CITIZEN: ("0F7A6E", "E3F3F0"),
    Family.CORE: ("154273", "E4ECF5"),
    Family.ADMIN: ("9A5B00", "FBEFDC"),
    Family.PLATFORM: ("5B4B8A", "ECE8F6"),
}
REGISTER_TITLE = "Bouwblokkenregister"
REGISTER_INTRO = (
    "Per bouwblok de verantwoordelijkheid, de modules waar het woont en de besluiten (ADR) die het "
    "vormgeven."
)
REGISTER_HEAD = ("ID", "Bouwblok", "Verantwoordelijkheid", "Modules", "ADR")


def plain(text: str) -> str:
    """Return ``text`` without draw.io markup and entities."""
    return html.unescape(re.sub(r"<[^>]+>", "", text.replace("<br>", " "))).strip()


def png_size(path: Path) -> tuple[int, int]:
    """Return (width, height) of a PNG from its IHDR chunk."""
    with path.open("rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError(f"{path} is not a PNG")
    width, height = struct.unpack(">II", head[16:24])
    return int(width), int(height)


def register_rows() -> list[tuple[Family, tuple[str, ...]]]:
    """Return the register as (family, cells) in plate order."""
    return [(b.family, (b.id, plain(b.name), b.responsibility, b.modules, b.adrs)) for b in BLOCKS]


def print_page(images: list[Path]) -> str:
    """Return the HTML page that the browser prints as the complete PDF."""
    plates = "".join(
        f'<section class="pg"><h2>Plaat {i} · {html.escape(p.title)}</h2>'
        f'<p class="m">{html.escape(p.description)}</p><img src="{img.as_uri()}" alt=""></section>'
        for i, (p, img) in enumerate(zip(PLATES, images, strict=True), start=1)
    )
    rows = "".join(
        f'<tr><td class="b {fam.value}"><span>{html.escape(c[0])}</span>{html.escape(c[1])}</td>'
        + "".join(f"<td>{html.escape(v)}</td>" for v in c[2:])
        + "</tr>"
        for fam, c in register_rows()
    )
    head = "".join(f"<th>{h}</th>" for h in ("Bouwblok", *REGISTER_HEAD[2:]))
    return f"""<!doctype html><html lang="nl"><meta charset="utf-8"><title>{TITLE}</title><style>
@page {{ size: A4 landscape; margin: 12mm 14mm; }}
body {{ font-family: Helvetica, Arial, sans-serif; color: #142033; font-size: 10pt; margin: 0; }}
h1 {{ color: #154273; font-size: 26pt; margin: 40mm 0 6mm; }}
h2 {{ color: #154273; font-size: 14pt; margin: 0 0 2mm; }}
.m {{ color: #5a6577; margin: 0 0 3mm; max-width: 240mm; }}
.pg {{ page-break-before: always; height: 183mm; display: flex; flex-direction: column; }}
.pg img {{ max-width: 100%; max-height: 165mm; object-fit: contain; align-self: center; }}
table {{ border-collapse: collapse; width: 100%; font-size: 8.5pt; }}
th, td {{ border: 1px solid #d5dbe4; padding: 3px 5px; text-align: left; vertical-align: top; }}
th {{ background: #154273; color: #fff; }}
td:nth-child(3) {{ font-family: Menlo, Consolas, monospace; font-size: 7.5pt; }}
td.b {{ font-weight: bold; white-space: nowrap; }}
td.b span {{ display: block; font-family: Menlo, monospace; font-size: 7.5pt; }}
.citizen {{ color: #0f7a6e; }} .core {{ color: #154273; }}
.admin {{ color: #9a5b00; }} .platform {{ color: #5b4b8a; }}
tr {{ page-break-inside: avoid; }}
</style>
<h1>{TITLE}</h1>
<p class="m" style="font-size:12pt">{html.escape(SUBTITLE)}</p>
<p>{html.escape(INTRO)}</p>
<p class="m">{html.escape(LEGEND_TEXT)}</p>
{plates}
<section style="page-break-before:always"><h2>{REGISTER_TITLE}</h2>
<p class="m">{REGISTER_INTRO}</p>
<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table></section></html>
"""


def _shade(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def write_docx(images: list[Path], out: Path) -> None:
    """Write the Word version (A4 landscape): intro, one plate per page, register table."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    accent, muted = RGBColor(0x15, 0x42, 0x73), RGBColor(0x5A, 0x65, 0x77)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21.0)
    section.left_margin = section.right_margin = Cm(1.6)
    section.top_margin = section.bottom_margin = Cm(1.4)
    usable = Cm(29.7 - 3.2)

    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    for style, size in (("Title", 26), ("Heading 1", 16)):
        doc.styles[style].font.name = "Calibri"
        doc.styles[style].font.size = Pt(size)
        doc.styles[style].font.color.rgb = accent

    def para(text: str, colour: Any = None, size: float | None = None) -> None:
        run = doc.add_paragraph().add_run(text)
        if colour is not None:
            run.font.color.rgb = colour
        if size is not None:
            run.font.size = Pt(size)

    doc.add_paragraph(TITLE, style="Title")
    para(SUBTITLE, muted)
    para(INTRO)
    para(LEGEND_TEXT, muted, 9.5)

    for i, (plate, image) in enumerate(zip(PLATES, images, strict=True), start=1):
        doc.add_page_break()
        doc.add_paragraph(f"Plaat {i} · {plate.title}", style="Heading 1")
        para(plate.description, muted, 9.5)
        width, height = png_size(image)
        doc.add_picture(str(image), width=min(usable, int(Cm(15.4) * width / height)))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()
    doc.add_paragraph(REGISTER_TITLE, style="Heading 1")
    para(REGISTER_INTRO, muted, 9.5)
    table = doc.add_table(rows=1, cols=len(REGISTER_HEAD))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, REGISTER_HEAD, strict=True):
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(cell, "154273")
    for family, values in register_rows():
        stroke, fill = FAMILY_HEX[family]
        cells = table.add_row().cells
        for j, (cell, text) in enumerate(zip(cells, values, strict=True)):
            run = cell.paragraphs[0].add_run(text)
            run.font.size = Pt(8 if j == 3 else 9)
            if j <= 1:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(stroke)
            if j == 3:
                run.font.name = "Consolas"
        _shade(cells[0], fill)
    widths = (1.2, 4.4, 10.6, 7.6, 2.6)
    for row in table.rows:
        for cell, width_cm in zip(row.cells, widths, strict=True):
            cell.width = Cm(width_cm)
    doc.save(str(out))
